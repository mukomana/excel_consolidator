"""
diagnose_report_dates.py

Run this against one of the actual failing files to see exactly what
extractor.py is finding (or not finding) at each step.

Usage:
    python diagnose_report_dates.py "SAPRIN_HDSS_WeeklyProgressReport V10.4 - 06 July 2026.xlsx"
"""

import sys
from openpyxl import load_workbook

from extractor import (
    COMMON_KPI_LABELS,
    find_label_cell,
    find_value_pair,
    get_workbook_report_period,
    extract_report,
    normalise,
)


def main(filepath):

    print("=" * 70)
    print(f"FILE: {filepath}")
    print("=" * 70)

    workbook = load_workbook(filepath, data_only=True)

    print("\nActual sheet names in this workbook:")
    for name in workbook.sheetnames:
        print(f"  - {name!r}")

    aliases = COMMON_KPI_LABELS["report_start_date"]
    print(f"\nSearching for label: {aliases}")

    for sheet_name in workbook.sheetnames:

        ws = workbook[sheet_name]

        label_cell = find_label_cell(ws, aliases)

        print(f"\n--- Sheet: {sheet_name!r} ---")

        if label_cell is None:
            print("  Label NOT found on this sheet.")
            continue

        print(f"  Label found at {label_cell.coordinate}: {label_cell.value!r}")

        # Show the raw values in the next few cells to the right, so we can
        # see exactly what's there (blank, merged, offset differently, etc.)
        for offset in range(1, 5):
            cell = ws.cell(row=label_cell.row, column=label_cell.column + offset)
            print(f"    +{offset} -> {cell.coordinate} = {cell.value!r}")

        start, end = find_value_pair(ws, aliases, offset_a=1, offset_b=2)
        print(f"  find_value_pair() result: start={start!r}, end={end!r}")

    print("\n--- get_workbook_report_period(workbook) ---")
    fallback = get_workbook_report_period(workbook)
    print(f"  Result: {fallback}")

    workbook.close()

    print("\n--- extract_report() full output ---")
    records = extract_report(filepath, "DIAGNOSTIC_SITE")
    for r in records:
        print(f"\n  survey_type={r.get('survey_type')}")
        print(f"    report_start_date = {r.get('report_start_date')!r}")
        print(f"    report_end_date   = {r.get('report_end_date')!r}")
        print(f"    _report_dates_inferred = {r.get('_report_dates_inferred')}")


if __name__ == "__main__":

    if len(sys.argv) < 2:
        print("Usage: python diagnose_report_dates.py <path_to_xlsx>")
        sys.exit(1)

    main(sys.argv[1])
