"""
extractor.py

Reads an IHS Weekly Report workbook and extracts KPIs.

Author: Freedom Mukomana
"""

from pathlib import Path
from typing import Dict, Optional

import re
from datetime import datetime

from openpyxl import load_workbook

###########################################################################
# KPI LABELS
###########################################################################

KPI_LABELS = {
    "start_date": [
        "Start Date"
    ],

    "end_date": [
        "End Date"
    ],

    "fca": [
        "FCA"
    ],

    "allocated": [
        "Individuals allocated"
    ],

    "completed": [
        "Completed"
    ],

    "completion_rate": [
        "Completion Rate"
    ],

    "contact_rate": [
        "Contact Rate"
    ],

    "dbs_rate": [
        "Participation Rate-DBS"
    ],

    "hiv_rate": [
        "Participation Rate-HIV"
    ],

    "height_weight_rate": [
        "Participation Rate-Height & Weight"
    ],

    "blood_glucose_rate": [
        "Participation Rate-Blood Glucose"
    ],

    "health_utilisation_rate": [
        "Participation Rate-Health Utilisation"
    ]
}

###############################################################################
# COMMON KPIs
###############################################################################

COMMON_KPI_LABELS = {

    "report_start_date": [
        "Which week are you reporting on (indicate start and end date)",
        "On which date did you start working on cases in this calendar year?"
    ],

    "report_end_date": [
        "Which week are you reporting on (indicate start and end date)",
        "On which date did you start working on cases in this calendar year?"
    ],

    "week": [
        "How many weeks have gone since starting this FCA/Trimester?"
    ],

    "time_elapsed_pct": [
        "What percentage of time have you now covered"
    ],

    "completion_rate": [
        "Completion Rate"
    ],

    "contact_rate": [
        "Contact  Rate"
    ],

    "participation_rate": [
        "Participation Rate"
    ]

}

###############################################################################
# FIELD + TELEPHONIC
###############################################################################

FIELD_TELEPHONE_LABELS = {

    "trimester": [
        "Which FCA /Trimester Are you currently working in?"
    ],

    "period_start": [
        "On which date did you start this 15 week FCA/Trimester?"
    ],

    "period_end": [
        "On which date did you start this 15 week FCA/Trimester?"
    ],

    "allocated": [
        "How many households are allocated to this FCA/Trimester?"
    ],

    "completed": [
        "How many households have been completed? (i.e. have been finalised, gone through QC and accepted in database as being completely surveyed)"
    ],

    "non_contact": [
        "Non-contact"
    ],

    "non_contact_rate": [
        "Non-contact"
    ],

    "passive_refusal": [
        "Passive Refusal"
    ],

    "passive_refusal_rate": [
        "Passive Refusal"
    ],

    "contacted": [
        "Contacted"
    ],

    "refused": [
        "Refused"
    ],

    "refusal_rate": [
        "Refusal Rate"
    ],

    "participated": [
        "Participated",
        "Particpated"
    ],

    "consented_yes": [
        "Consented YES"
    ],

    "consented_no": [
        "Consented No"
    ]

}

###############################################################################
# VERBAL AUTOPSY
###############################################################################

VA_LABELS = {

    "reporting_year": [
        "Which year you currently working in?"
    ],

    "year_start_date": [
        "On which date did you start working on cases in this calendar year?"
    ],

    "year_end_date": [
        "On which date did you start working on cases in this calendar year?"
    ],

    "allocated": [
        "How many VA cases are allocated to this calendar year?"
    ],

    "completed": [
        "How many VA cases have been completed? (i.e. have been finalised, gone through QC and accepted in database as being completely surveyed)"
    ],

    "premature": [
        "Premature"
    ],

    "premature_rate": [
        "Premature"
    ],

    "non_contact": [
        "Non-contact"
    ],

    "non_contact_rate": [
        "Non-contact"
    ],

    "contacted": [
        "Contacted"
    ],

    "refused": [
        "Refused"
    ],

    "participated": [
        "Participated",
        "Particpated"
    ],

    "consented_yes": [
        "Yes"
    ],

    "consented_no": [
        "No"
    ]

}

###############################################################################
# PAIRED FIELDS
#
# Some rows in the source sheet carry TWO related values on the same row
# (a start value in one column and an end value further along), but both
# logical fields share the exact same row label. find_value() can only ever
# return one value for a given label, so if start_date and end_date fields
# both search for the same label text, the second call just re-finds the
# same cell and returns the same (start) value again -- end_date silently
# gets overwritten with start_date.
#
# This map tells extract_sheet() which field-pairs need to be read together,
# from a single row, at two different column offsets.
###############################################################################

PAIRED_FIELDS = {
    ("report_start_date", "report_end_date"): (
        ["Which week are you reporting on (indicate start and end date)"],
        1, 2  # value at label_col+1 -> start, label_col+2 -> end
    ),
    ("period_start", "period_end"): (
        ["On which date did you start this 15 week FCA/Trimester?"],
        1, 2
    ),
    ("year_start_date", "year_end_date"): (
        ["On which date did you start working on cases in this calendar year?"],
        1, 2
    ),
}

###########################################################################
# HELPERS
###########################################################################

def normalise(value):

    if value is None:
        return ""

    value = str(value)

    value = value.replace("\n", " ")

    value = re.sub(r"\s+", " ", value)

    return value.strip().lower()


def percentage_to_float(value):

    if value is None:
        return None

    if isinstance(value, (int, float)):
        if value <= 1:
            return value * 100
        return value

    value = str(value).replace("%", "").strip()

    try:
        return float(value)
    except Exception:
        return value


def parse_date(value):

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    return value


###########################################################################
# MERGED-CELL HELPER
#
# In openpyxl, only the TOP-LEFT cell of a merged range actually stores a
# value -- every other cell in that range reads as None. These report
# templates use merged cells for several value cells, so a raw
# ws.cell(row, col).value can silently come back empty even when a value
# is visibly there. resolve_merge() redirects any cell inside a merge to
# its top-left cell before reading.
###########################################################################

def resolve_merge(ws, cell):

    for merged_range in ws.merged_cells.ranges:

        if cell.coordinate in merged_range:

            return ws.cell(
                row=merged_range.min_row,
                column=merged_range.min_col
            )

    return cell


def get_cell_value(ws, row, column):

    cell = ws.cell(row=row, column=column)

    cell = resolve_merge(ws, cell)

    return cell.value


###########################################################################
# FIND LABEL
###########################################################################

def find_label_cell(ws, labels):

    """
    Search worksheet for one of the labels. Returns the matching cell,
    or None.
    """

    labels = [normalise(x) for x in labels]

    for row in ws.iter_rows():

        for cell in row:

            if normalise(cell.value) in labels:

                return cell

    return None


def find_value(ws, labels):

    """
    Search worksheet for one of the labels.

    Returns the value immediately to the right (merge-aware).
    """

    label_cell = find_label_cell(ws, labels)

    if label_cell is None:
        return None

    value = get_cell_value(ws, label_cell.row, label_cell.column + 1)

    if value is None:

        # Try two columns to the right

        value = get_cell_value(ws, label_cell.row, label_cell.column + 2)

    return value


def find_value_pair(ws, labels, offset_a=1, offset_b=2):

    """
    Search worksheet for one of the labels, and return TWO values from
    that same row (merge-aware) -- e.g. a start value and an end value
    that live on the same labelled row but in different columns.

    Returns (value_a, value_b), or (None, None) if the label isn't found.
    """

    label_cell = find_label_cell(ws, labels)

    if label_cell is None:
        return None, None

    value_a = get_cell_value(ws, label_cell.row, label_cell.column + offset_a)
    value_b = get_cell_value(ws, label_cell.row, label_cell.column + offset_b)

    return value_a, value_b


def extract_sheet(sheet, survey_type, labels):

    record = {

        "survey_type": survey_type

    }

    ###########################################################################
    # Paired fields (start/end values sharing one row label)
    #
    # Handled first, and excluded from the generic single-value loops below
    # so they don't get re-processed (and re-overwritten) a second time.
    ###########################################################################

    paired_field_names = set()

    for (field_a, field_b), (aliases, offset_a, offset_b) in PAIRED_FIELDS.items():

        # Only apply a pair if both fields are actually relevant to this
        # sheet (i.e. present in either the common or survey-specific labels)
        relevant = (
            field_a in COMMON_KPI_LABELS or field_a in labels
        ) and (
            field_b in COMMON_KPI_LABELS or field_b in labels
        )

        if not relevant:
            continue

        value_a, value_b = find_value_pair(sheet, aliases, offset_a, offset_b)

        if "date" in field_a:
            value_a = parse_date(value_a)
        if "date" in field_b:
            value_b = parse_date(value_b)

        if value_a is not None:
            record[field_a] = value_a
        if value_b is not None:
            record[field_b] = value_b

        paired_field_names.add(field_a)
        paired_field_names.add(field_b)

    ###########################################################################
    # Common KPIs
    ###########################################################################

    for field, aliases in COMMON_KPI_LABELS.items():

        if field in paired_field_names:
            continue

        value = find_value(sheet, aliases)

        if value is None:
            continue

        if "rate" in field:
            value = percentage_to_float(value)

        if "date" in field:
            value = parse_date(value)

        record[field] = value

    ###########################################################################
    # Survey-specific KPIs
    ###########################################################################

    for field, aliases in labels.items():

        if field in paired_field_names:
            continue

        value = find_value(sheet, aliases)

        if value is None:
            continue

        if "rate" in field:
            value = percentage_to_float(value)

        if "date" in field:
            value = parse_date(value)

        record[field] = value

    return record

###########################################################################
# EXTRACT
###########################################################################

def extract_report(filepath, site_name):

    workbook = load_workbook(
        filepath,
        data_only=True
    )

    records = []

    ###########################################################################
    # Field Household Surveillance
    ###########################################################################

    if "Field-HouseholdSurveillance" in workbook.sheetnames:

        record = {

            "site": site_name,

            "source_file": Path(filepath).name

        }

        record.update(

            extract_sheet(

                workbook["Field-HouseholdSurveillance"],

                "Field",

                FIELD_TELEPHONE_LABELS

            )

        )

        records.append(record)

    ###########################################################################
    # Telephonic Household Surveillance
    ###########################################################################

    if "Telephonic HouseholdSurveillance" in workbook.sheetnames:

        record = {

            "site": site_name,

            "source_file": Path(filepath).name

        }

        record.update(

            extract_sheet(

                workbook["Telephonic HouseholdSurveillance"],

                "Telephonic",

                FIELD_TELEPHONE_LABELS

            )

        )

        records.append(record)

    ###########################################################################
    # Verbal Autopsy
    ###########################################################################

    if "VerbalAutopsy" in workbook.sheetnames:

        record = {

            "site": site_name,

            "source_file": Path(filepath).name

        }

        record.update(

            extract_sheet(

                workbook["VerbalAutopsy"],

                "Verbal Autopsy",

                VA_LABELS

            )

        )

        records.append(record)

    workbook.close()

    return records

###########################################################################
# VALIDATION
###########################################################################

REQUIRED_FIELDS = [

    "start_date",

    "end_date",

    "allocated",

    "completed"

]


def validate_record(record):

    missing = []

    for field in REQUIRED_FIELDS:

        if field not in record:

            missing.append(field)

    return missing


###########################################################################
# TEST
###########################################################################

if __name__ == "__main__":

    report = extract_report(

        "Sample Weekly Report.xlsx",

        "Johannesburg"

    )

    print(report)

    print()

    print(validate_record(report))